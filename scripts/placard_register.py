"""The register workbook that travels with a set of placard workbooks.

The placard workbooks are for printing, so they hold nothing else. This
workbook is for running the program afterwards: one row per machine, the
energy steps behind every placard, the findings of the pre-delivery audit,
and the annual periodic-inspection log OSHA 1910.147(c)(6) requires.

It is built to keep working after it leaves us. Every count and status is a
formula over the sheets beside it, so when the client logs an inspection or
closes a finding, the register and dashboard update themselves:

    Inspection Log  ─►  Register (last inspection, next due, status)
    Audit Findings  ─►  Register (open findings per machine)
    Register        ─►  Dashboard (per-department roll-up)
"""

from __future__ import annotations

import os
import urllib.parse
from datetime import date, datetime
from pathlib import Path

from export_placards_xlsx import (
    ENERGY_CODES,
    FONT_NAME,
    NAVY_HEADER,
    PHASE_LABELS,
    SLATE_TEXT,
    WHITE,
    PlacardPage,
    _grouped_steps,
    clean,
    group_by_department,
    placard_text,
)
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

LOG_FIRST_ROW = 8        # first fillable row of the Inspection Log
LOG_ROWS = 2000          # fillable rows — four years of a ~500-machine site
INPUT_FILL = "FFF8D6"    # pale yellow: the cells the client is meant to edit
EXAMPLE_GREY = "8C8C99"
FINDING_STATUSES = ("Open", "In progress", "Corrected", "Not applicable")
OPEN_STATUSES = ("Open", "In progress")
SEVERITY_STYLE = {          # fill, text
    "CRITICAL": ("8B0A1A", WHITE),
    "HIGH": ("E07B00", WHITE),
    "MEDIUM": ("FFD900", "1A1A1A"),
    "LOW": ("D1D9E6", "1A1A1A"),
}
SEVERITY_ORDER = {name: rank for rank, name in enumerate(SEVERITY_STYLE)}

THIN = Side(style="thin", color="D1D9E6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _font(size=10, bold=False, color=SLATE_TEXT, italic=False, underline=None):
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic,
                underline=underline)


def _header(ws, row: int, labels: list[str], widths: list[float]) -> None:
    for column, (label, width) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(row=row, column=column, value=label)
        cell.font = _font(10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY_HEADER)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.row_dimensions[row].height = 32


def _body(cell, wrap=False, fmt=None, fill=None, bold=False, color=SLATE_TEXT):
    cell.font = _font(10, bold=bold, color=color)
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    cell.border = BOX
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _link_to(cell, workbook_path: str, sheet: str, label: str) -> None:
    """Link a cell to a sheet in another workbook in the same folder tree.

    The target is relative, so the link survives the folder being copied,
    zipped or moved to a shared drive — as long as it moves as a whole.
    """
    location = "'{}'!A1".format(sheet.replace("'", "''"))
    cell.value = label
    cell.hyperlink = Hyperlink(ref=cell.coordinate,
                               target=urllib.parse.quote(workbook_path),
                               location=location)
    cell.font = _font(10, color="1155CC", underline="single")


def _parse_date(value) -> date | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).date()
    except ValueError:
        return None


# ── Sheets ──────────────────────────────────────────────────────────────────
def _energy_codes_sheet(ws) -> None:
    _header(ws, 1, ["Code", "Energy type (English)", "Tipo de energía (Español)"],
            [10, 26, 28])
    for code, label_en, label_es, fill, text in ENERGY_CODES:
        ws.append([code, label_en, label_es])
        row = ws.max_row
        for column in (1, 2, 3):
            _body(ws.cell(row=row, column=column))
        badge = ws.cell(row=row, column=1)
        badge.fill = PatternFill("solid", fgColor=fill)
        badge.font = _font(10, bold=True, color=text)
    ws.append([])
    ws.append(["N is not an energy source: it marks the shutdown, lock-application and "
               "try-out rows that apply to the whole machine."])
    ws.cell(row=ws.max_row, column=1).font = _font(9, italic=True, color=EXAMPLE_GREY)
    ws.freeze_panes = "A2"


def _audit_sheet(ws, findings: list[dict], equipment_by_id: dict, department_of: dict,
                 link_for: dict) -> int:
    labels = ["#", "Severity", "Equipment ID", "Description", "Department",
              "Energy code", "Procedure phase", "Field", "Issue", "Detail",
              "Recommended correction", "Found by", "Placard",
              "Status", "Owner", "Date corrected", "Notes"]
    widths = [6, 11, 16, 34, 20, 9, 16, 22, 30, 60, 52, 18, 18, 14, 16, 13, 30]
    _header(ws, 1, labels, widths)
    ordered = sorted(findings, key=lambda f: (SEVERITY_ORDER.get(f["severity"], 9),
                                              f["equipment_id"], f.get("step_number") or 0))
    for number, finding in enumerate(ordered, 1):
        equipment_id = finding["equipment_id"]
        equipment = equipment_by_id.get(equipment_id, {})
        phase = PHASE_LABELS["en"].get(finding.get("step_type") or "", "")
        ws.append([number, finding["severity"], equipment_id,
                   clean(equipment.get("description")), department_of.get(equipment_id, ""),
                   finding.get("energy_type") or "", phase.title(),
                   finding.get("field") or "", finding["check"], finding["detail"],
                   finding["action"], finding.get("source") or "Automated rule check",
                   None, "Open", None, None, None])
        row = ws.max_row
        for column in range(1, len(labels) + 1):
            _body(ws.cell(row=row, column=column), wrap=column in (4, 9, 10, 11, 17))
        fill, text = SEVERITY_STYLE.get(finding["severity"], ("FFFFFF", SLATE_TEXT))
        ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=row, column=2).font = _font(10, bold=True, color=text)
        if equipment_id in link_for:
            _link_to(ws.cell(row=row, column=13), *link_for[equipment_id],
                     label="Open placard")
        for column in (14, 15, 16, 17):
            ws.cell(row=row, column=column).fill = PatternFill("solid", fgColor=INPUT_FILL)
        ws.cell(row=row, column=16).number_format = "yyyy-mm-dd"

    last = max(ws.max_row, 2)
    status = DataValidation(type="list", formula1='"{}"'.format(",".join(FINDING_STATUSES)),
                            allow_blank=False, showErrorMessage=True,
                            errorTitle="Status", error="Pick a status from the list.")
    ws.add_data_validation(status)
    status.add(f"N2:N{last}")
    corrected = DataValidation(type="date", operator="greaterThan", formula1="36526",
                               showErrorMessage=True, error="Enter a date (yyyy-mm-dd).")
    ws.add_data_validation(corrected)
    corrected.add(f"P2:P{last}")
    # A finding marked corrected greys out, so what is left open stands out.
    ws.conditional_formatting.add(
        f"A2:Q{last}",
        FormulaRule(formula=['OR($N2="Corrected",$N2="Not applicable")'],
                    font=Font(color=EXAMPLE_GREY, strike=True)))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:Q{last}"
    return len(ordered)


def _steps_sheet(ws, equipment: list[dict], steps_by_id: dict, department_of: dict) -> None:
    labels = ["Equipment ID", "Department", "Print order", "Procedure phase", "Energy code",
              "Energy type", "Try-out required",
              "Energy tag & description", "Isolation procedure & lockout devices",
              "Method of verification",
              "Etiqueta y descripción (ES)", "Procedimiento de aislamiento (ES)",
              "Método de verificación (ES)", "Open findings (this machine)"]
    widths = [16, 20, 8, 18, 8, 16, 9, 40, 60, 46, 40, 60, 46, 12]
    _header(ws, 1, labels, widths)
    for machine in sorted(equipment, key=lambda e: e["equipment_id"]):
        equipment_id = machine["equipment_id"]
        for order, step in enumerate(_grouped_steps(steps_by_id.get(equipment_id, [])), 1):
            ws.append([equipment_id, department_of[equipment_id], order,
                       PHASE_LABELS["en"].get(step.get("step_type"),
                                              clean(step.get("step_type"))).title(),
                       clean(step.get("energy_type")), None,
                       "Yes" if step.get("tryout_required") else "No",
                       placard_text(step.get("tag_description")),
                       placard_text(step.get("isolation_procedure")),
                       placard_text(step.get("method_of_verification")),
                       placard_text(step.get("tag_description_es")),
                       placard_text(step.get("isolation_procedure_es")),
                       placard_text(step.get("method_of_verification_es")), None])
            row = ws.max_row
            ws.cell(row=row, column=6).value = (
                f"=IFERROR(INDEX('Energy Codes'!$B$2:$B$40,"
                f"MATCH(E{row},'Energy Codes'!$A$2:$A$40,0)),\"Unknown code\")")
            ws.cell(row=row, column=14).value = (
                f"=COUNTIFS('Audit Findings'!$C:$C,$A{row},'Audit Findings'!$N:$N,\"Open\")"
                f"+COUNTIFS('Audit Findings'!$C:$C,$A{row},'Audit Findings'!$N:$N,\"In progress\")")
            for column in range(1, len(labels) + 1):
                _body(ws.cell(row=row, column=column), wrap=column in range(8, 14))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{max(ws.max_row, 2)}"


def _inspection_log_sheet(ws, register_rows: int) -> None:
    ws["A1"] = "Periodic Inspection Log — Registro de Inspección Periódica"
    ws["A1"].font = _font(14, bold=True, color=NAVY_HEADER)
    notes = [
        "OSHA 29 CFR 1910.147(c)(6) / Cal/OSHA T8 §3314(j): inspect every energy-control "
        "procedure at least once a year. The inspector must be an authorized employee other "
        "than the one(s) using the procedure, and the record must name the machine, the date, "
        "the employees observed and the inspector.",
        "Fill in the yellow columns — one row per inspection. Grey columns fill themselves. "
        "The Register and Dashboard read this sheet: each machine's last inspection, next due "
        "date and OVERDUE status update as you log here.",
        "Llene las columnas amarillas, una fila por inspección. El Registro y el Tablero se "
        "actualizan automáticamente.",
    ]
    for offset, text in enumerate(notes, 2):
        ws.cell(row=offset, column=1, value=text).font = _font(9, color=SLATE_TEXT)
        ws.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=10)
        ws.cell(row=offset, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[offset].height = 28

    labels = ["Inspection date", "Equipment ID", "Description (auto)", "Department (auto)",
              "Inspector (authorized employee)", "Employees observed / reviewed",
              "Procedure adequate?", "Deviations or inadequacies found",
              "Corrective action taken", "Inspector signature / initials"]
    widths = [14, 16, 36, 20, 24, 30, 14, 40, 40, 22]
    header_row = LOG_FIRST_ROW - 2
    _header(ws, header_row, labels, widths)

    example_row = LOG_FIRST_ROW - 1
    example = [date(2026, 11, 2), "EXAMPLE", "Example row — not counted",
               "", "R. Martinez (maint. lead)", "J. Lopez, A. Chen",
               "Needs revision", "Pneumatic step did not say where the bleed valve is",
               "Added valve location to step; re-trained J. Lopez", "RM"]
    for column, value in enumerate(example, 1):
        cell = ws.cell(row=example_row, column=column, value=value)
        cell.font = _font(9, italic=True, color=EXAMPLE_GREY)
        cell.border = BOX
    ws.cell(row=example_row, column=1).number_format = "yyyy-mm-dd"

    last = LOG_FIRST_ROW + LOG_ROWS - 1
    for row in range(LOG_FIRST_ROW, last + 1):
        for column in range(1, 11):
            cell = ws.cell(row=row, column=column)
            cell.border = BOX
            cell.font = _font(10)
            if column in (3, 4):
                cell.fill = PatternFill("solid", fgColor="F5F7FA")
            else:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
        ws.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=3).value = (
            f'=IF($B{row}="","",IFERROR(INDEX(Register!$B$2:$B${register_rows + 1},'
            f'MATCH($B{row},Register!$A$2:$A${register_rows + 1},0)),"Not in register"))')
        ws.cell(row=row, column=4).value = (
            f'=IF($B{row}="","",IFERROR(INDEX(Register!$C$2:$C${register_rows + 1},'
            f'MATCH($B{row},Register!$A$2:$A${register_rows + 1},0)),""))')

    when = DataValidation(type="date", operator="greaterThan", formula1="36526",
                          showErrorMessage=True, error="Enter a date (yyyy-mm-dd).")
    machine = DataValidation(type="list",
                             formula1=f"=Register!$A$2:$A${register_rows + 1}",
                             showErrorMessage=True, errorTitle="Equipment ID",
                             error="Pick an Equipment ID from the Register.")
    verdict = DataValidation(type="list", formula1='"Yes,Needs revision"', allow_blank=True)
    for validation, column in ((when, "A"), (machine, "B"), (verdict, "G")):
        ws.add_data_validation(validation)
        validation.add(f"{column}{LOG_FIRST_ROW}:{column}{last}")
    ws.freeze_panes = f"C{LOG_FIRST_ROW}"


def _register_sheet(ws, equipment: list[dict], steps_by_id: dict, department_of: dict,
                    photos: Path, link_for: dict) -> int:
    labels = ["Equipment ID", "Description", "Department", "Energy sources",
              "Energy steps", "Photos on placard (of 2)", "Procedure verified",
              "Verified on", "Spanish reviewed", "Last periodic inspection",
              "Next inspection due", "Inspection status",
              "Open CRITICAL findings", "Open HIGH findings", "Placard (EN / ES)", "Notes"]
    widths = [16, 40, 20, 14, 9, 11, 11, 12, 10, 14, 14, 18, 11, 11, 16, 30]
    _header(ws, 1, labels, widths)

    log = "'Inspection Log'"
    log_dates = f"{log}!$A${LOG_FIRST_ROW}:$A${LOG_FIRST_ROW + LOG_ROWS - 1}"
    log_ids = f"{log}!$B${LOG_FIRST_ROW}:$B${LOG_FIRST_ROW + LOG_ROWS - 1}"
    audit = "'Audit Findings'"

    for machine in sorted(equipment, key=lambda e: e["equipment_id"]):
        equipment_id = machine["equipment_id"]
        steps = steps_by_id.get(equipment_id, [])
        codes = sorted({clean(s.get("energy_type")) for s in steps} - {"N", ""})
        stem = equipment_id.replace("/", "_")
        photo_count = sum((photos / f"{stem}_{kind}.jpg").exists() for kind in ("equip", "iso"))
        ws.append([equipment_id, clean(machine.get("description")), department_of[equipment_id],
                   " · ".join(codes), len(steps), photo_count,
                   "Yes" if machine.get("verified") else "No",
                   _parse_date(machine.get("verified_date")),
                   "Yes" if machine.get("spanish_reviewed") else "No"])
        row = ws.max_row
        ws.cell(row=row, column=10).value = (
            f'=IF(COUNTIF({log_ids},$A{row})=0,"",_xlfn.MAXIFS({log_dates},{log_ids},$A{row}))')
        ws.cell(row=row, column=11).value = f'=IF($J{row}="","",EDATE($J{row},12))'
        ws.cell(row=row, column=12).value = (
            f'=IF($J{row}="","Not yet recorded",IF($K{row}<TODAY(),"OVERDUE",'
            f'IF($K{row}-TODAY()<=30,"Due within 30 days","Current")))')
        for column, severity in ((13, "CRITICAL"), (14, "HIGH")):
            ws.cell(row=row, column=column).value = " + ".join(
                f"COUNTIFS({audit}!$C:$C,$A{row},{audit}!$B:$B,\"{severity}\","
                f"{audit}!$N:$N,\"{status}\")" for status in OPEN_STATUSES
            ).join(("=", ""))
        for column in range(1, len(labels) + 1):
            _body(ws.cell(row=row, column=column), wrap=column == 2)
        for column in (8, 10, 11):
            ws.cell(row=row, column=column).number_format = "yyyy-mm-dd"
        ws.cell(row=row, column=16).fill = PatternFill("solid", fgColor=INPUT_FILL)
        if equipment_id in link_for:
            _link_to(ws.cell(row=row, column=15), *link_for[equipment_id],
                     label="Open placard")

    last = max(ws.max_row, 2)
    status_range = f"L2:L{last}"
    for text, fill, color in (("OVERDUE", "8B0A1A", WHITE),
                              ("Due within 30 days", "FFD900", "1A1A1A"),
                              ("Current", "D9F2DC", "1A5C24"),
                              ("Not yet recorded", "EEEEF2", "555566")):
        ws.conditional_formatting.add(status_range, CellIsRule(
            operator="equal", formula=[f'"{text}"'],
            fill=PatternFill("solid", fgColor=fill), font=Font(color=color, bold=True)))
    for column, fill in (("M", "8B0A1A"), ("N", "E07B00")):
        ws.conditional_formatting.add(f"{column}2:{column}{last}", CellIsRule(
            operator="greaterThan", formula=["0"],
            fill=PatternFill("solid", fgColor=fill), font=Font(color=WHITE, bold=True)))
    ws.conditional_formatting.add(f"F2:F{last}", CellIsRule(
        operator="lessThan", formula=["2"], font=Font(color="BF1414", bold=True)))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:P{last}"
    return len(equipment)


def _dashboard_sheet(ws, departments: list[str], machines: int, generated: str) -> None:
    ws["A1"] = "LOTO Program Dashboard — Tablero del Programa LOTO"
    ws["A1"].font = _font(15, bold=True, color=NAVY_HEADER)
    ws["A2"] = (f"Placards generated {generated}. Every figure below is a live formula over the "
                "Register, Audit Findings and Inspection Log sheets.")
    ws["A2"].font = _font(9, italic=True, color=EXAMPLE_GREY)

    reg = f"Register!$A$2:$A${machines + 1}"
    col = lambda letter: f"Register!${letter}$2:${letter}${machines + 1}"
    audit_open = lambda severity: " + ".join(
        f"COUNTIFS('Audit Findings'!$B:$B,\"{severity}\",'Audit Findings'!$N:$N,\"{s}\")"
        for s in OPEN_STATUSES)
    kpis = [
        ("Machines with a placard", f"=COUNTA({reg})", "0"),
        ("Energy isolation steps", f"=SUM({col('E')})", "0"),
        ("Placard photos embedded", f"=SUM({col('F')})/(2*COUNTA({reg}))", "0.0%"),
        ("Procedures verified", f"=COUNTIF({col('G')},\"Yes\")", "0"),
        ("Spanish translations reviewed", f"=COUNTIF({col('I')},\"Yes\")", "0"),
        ("Open CRITICAL audit findings (incl. site-wide)", f"={audit_open('CRITICAL')}", "0"),
        ("Open HIGH audit findings (incl. site-wide)", f"={audit_open('HIGH')}", "0"),
        ("Periodic inspections OVERDUE", f"=COUNTIF({col('L')},\"OVERDUE\")", "0"),
        ("Periodic inspections not yet recorded",
         f"=COUNTIF({col('L')},\"Not yet recorded\")", "0"),
    ]
    for offset, (label, formula, fmt) in enumerate(kpis, 4):
        ws.cell(row=offset, column=1, value=label).font = _font(10)
        value = ws.cell(row=offset, column=2, value=formula)
        value.font = _font(11, bold=True, color=NAVY_HEADER)
        value.number_format = fmt
        value.alignment = Alignment(horizontal="right")
    for row in (9, 10):
        ws.conditional_formatting.add(f"B{row}", CellIsRule(
            operator="greaterThan", formula=["0"], font=Font(color="BF1414", bold=True)))

    header_row = 4 + len(kpis) + 1
    labels = ["Department", "Placards", "Energy steps", "Photos missing",
              "Procedures verified", "Open CRITICAL", "Open HIGH",
              "Inspections OVERDUE", "Inspections not recorded"]
    _header(ws, header_row, labels, [46, 11, 11, 11, 12, 11, 11, 13, 14])
    first = header_row + 1
    for department in departments:
        ws.append([department])
        row = ws.max_row
        dept = f"{col('C')},$A{row}"
        formulas = [
            f"=COUNTIF({dept})",
            f"=SUMIF({dept},{col('E')})",
            f"=COUNTIFS({dept},{col('F')},\"<2\")",
            f"=COUNTIFS({dept},{col('G')},\"Yes\")",
            f"=SUMIF({dept},{col('M')})",
            f"=SUMIF({dept},{col('N')})",
            f"=COUNTIFS({dept},{col('L')},\"OVERDUE\")",
            f"=COUNTIFS({dept},{col('L')},\"Not yet recorded\")",
        ]
        for column, formula in enumerate(formulas, 2):
            ws.cell(row=row, column=column, value=formula)
        for column in range(1, len(labels) + 1):
            _body(ws.cell(row=row, column=column))
    total = ws.max_row + 1
    ws.cell(row=total, column=1, value="Total")
    for column in range(2, len(labels) + 1):
        letter = get_column_letter(column)
        ws.cell(row=total, column=column, value=f"=SUM({letter}{first}:{letter}{total - 1})")
    for column in range(1, len(labels) + 1):
        _body(ws.cell(row=total, column=column), bold=True, fill="EEF2F8")
    note = ws.cell(row=total + 1, column=1, value=(
        "Department rows count findings on that department's machines. Site-wide findings "
        "(Equipment ID 'SITE-WIDE' on the Audit Findings sheet) are in the totals above only."))
    note.font = _font(9, italic=True, color=EXAMPLE_GREY)
    for letter, fill in (("F", "8B0A1A"), ("G", "E07B00"), ("H", "8B0A1A")):
        ws.conditional_formatting.add(f"{letter}{first}:{letter}{total - 1}", CellIsRule(
            operator="greaterThan", formula=["0"],
            fill=PatternFill("solid", fgColor=fill), font=Font(color=WHITE, bold=True)))
    ws.freeze_panes = f"A{first}"


DRAFT_STATUSES = ("Draft — verify on site", "Field-verified", "Applied to placard", "Rejected")


def _draft_corrections_sheet(ws, proposed: list[dict], department_of: dict,
                             link_for: dict) -> None:
    """Replacement steps drafted by the pre-delivery review.

    These are starting points for the walkdown, never procedures: every
    value the data could not supply is bracketed [VERIFY ON SITE], and the
    status column exists so nothing reaches a placard unverified.
    """
    ws["A1"] = ("DRAFT corrections — not approved procedures. Verify every step at the "
                "machine before it is applied to a placard.")
    ws["A1"].font = _font(12, bold=True, color="BF1414")
    labels = ["Equipment ID", "Department", "Energy code", "Replaces",
              "Energy tag & description", "Isolation procedure & lockout devices",
              "Method of verification", "Must be confirmed on site", "Why this step",
              "Placard", "Status", "Verified by", "Date"]
    widths = [16, 20, 8, 14, 44, 70, 52, 40, 36, 14, 20, 18, 12]
    _header(ws, 3, labels, widths)
    for draft in sorted(proposed, key=lambda d: d["equipment_id"]):
        replaces = ""
        if draft.get("replaces_energy_type") or draft.get("replaces_step_number"):
            replaces = (f"{draft.get('replaces_energy_type') or ''} step "
                        f"{draft.get('replaces_step_number') or ''}").strip()
        ws.append([draft["equipment_id"], department_of.get(draft["equipment_id"], ""),
                   draft["energy_type"], replaces or "New step",
                   placard_text(draft["tag_description"]),
                   placard_text(draft["isolation_procedure"]),
                   placard_text(draft["method_of_verification"]),
                   "\n".join(f"• {item}" for item in draft.get("needs_site_verification") or []),
                   draft.get("rationale") or "", None, DRAFT_STATUSES[0], None, None])
        row = ws.max_row
        for column in range(1, len(labels) + 1):
            _body(ws.cell(row=row, column=column), wrap=column in (5, 6, 7, 8, 9))
        for column in (11, 12, 13):
            ws.cell(row=row, column=column).fill = PatternFill("solid", fgColor=INPUT_FILL)
        ws.cell(row=row, column=13).number_format = "yyyy-mm-dd"
        if draft["equipment_id"] in link_for:
            _link_to(ws.cell(row=row, column=10), *link_for[draft["equipment_id"]],
                     label="Open placard")
    last = max(ws.max_row, 4)
    status = DataValidation(type="list", formula1='"{}"'.format(",".join(DRAFT_STATUSES)),
                            showErrorMessage=True, error="Pick a status from the list.")
    ws.add_data_validation(status)
    status.add(f"K4:K{last}")
    ws.freeze_panes = "B4"
    ws.auto_filter.ref = f"A3:M{last}"


def _readme_sheet(ws, lines: list[tuple[str, str]]) -> None:
    """`lines` is (style, text); style is title, heading, body, note or blank."""
    styles = {"title": _font(16, bold=True, color=NAVY_HEADER),
              "heading": _font(12, bold=True, color=NAVY_HEADER),
              "body": _font(10), "note": _font(9, italic=True, color=EXAMPLE_GREY),
              "blank": _font(10)}
    for style, text in lines:
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = styles[style]
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    ws.sheet_view.showGridLines = False


def build_register(equipment: list[dict], steps: list[dict], pages: list[PlacardPage],
                   findings: list[dict], photos: Path, generated: str, out_path: Path,
                   placard_dir: Path, readme: list[tuple[str, str]],
                   proposed: list[dict] | None = None) -> None:
    """Write the register. `placard_dir` is where the placard workbooks sit,
    so links can be made relative to the register's own folder."""
    steps_by_id: dict[str, list[dict]] = {}
    for step in steps:
        steps_by_id.setdefault(step["equipment_id"], []).append(step)
    groups = group_by_department(equipment)
    department_of = {row["equipment_id"]: label for label, rows in groups for row in rows}
    equipment_by_id = {row["equipment_id"]: row for row in equipment}

    relative = os.path.relpath(placard_dir, out_path.parent)
    link_for = {}
    for page in pages:
        if page.language == "en":
            link_for[page.equipment_id] = (
                Path(relative, page.workbook).as_posix(), page.sheet)

    workbook = Workbook()
    readme_ws = workbook.active
    readme_ws.title = "Read Me"
    dashboard, register, log, audit, steps_ws, codes = (
        workbook.create_sheet(name) for name in
        ("Dashboard", "Register", "Inspection Log", "Audit Findings", "Energy Steps",
         "Energy Codes"))

    _readme_sheet(readme_ws, readme)
    rows = _register_sheet(register, equipment, steps_by_id, department_of, photos, link_for)
    _dashboard_sheet(dashboard, [label for label, _ in groups], rows, generated)
    _inspection_log_sheet(log, rows)
    _audit_sheet(audit, findings, equipment_by_id, department_of, link_for)
    _steps_sheet(steps_ws, equipment, steps_by_id, department_of)
    _energy_codes_sheet(codes)
    if proposed:
        drafts = workbook.create_sheet("Draft Corrections", workbook.index(steps_ws))
        _draft_corrections_sheet(drafts, proposed, department_of, link_for)

    for ws in workbook.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1" if ws.title in ("Register", "Audit Findings",
                                                     "Energy Steps") else None
    workbook.calculation.fullCalcOnLoad = True
    workbook.active = workbook.index(dashboard)
    workbook.save(out_path)
